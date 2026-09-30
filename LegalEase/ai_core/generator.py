"""Export a drafted document as TXT, DOCX or PDF (with logo, terms table, footer)."""
import io
from pathlib import Path

import config
from ai_core.text_utils import parse_blocks, xml_escape


def _has_non_latin(text: str) -> bool:
    return any(ord(c) > 255 for c in text)


def custom_font_available() -> bool:
    return Path(config.FONT_PATH).exists()


# ---------------------------------------------------------------- TXT
def build_txt(title, body, terms) -> bytes:
    out = [title.upper(), "=" * len(title), ""]
    if terms:
        out.append("KEY TERMS")
        out += [f"  {k}: {v}" for k, v in terms]
        out.append("")
    out += [body, "", "-" * 60, config.DISCLAIMER, config.COMPANY_FOOTER]
    return "\n".join(out).encode("utf-8")


# ---------------------------------------------------------------- DOCX
def build_docx(title, body, terms, include_logo=True) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if include_logo and Path(config.LOGO_PATH).exists():
        doc.add_picture(str(config.LOGO_PATH), width=Inches(1.6))
    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if terms:
        doc.add_heading("Key Terms", level=2)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for k, v in terms:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = k, v
            for run in cells[0].paragraphs[0].runs:
                run.bold = True
        doc.add_paragraph()

    for kind, text in parse_blocks(body):
        if kind == "h1":
            continue  # title already shown
        if kind == "h2":
            doc.add_heading(text, level=2)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)

    note = doc.add_paragraph()
    run = note.add_run(config.DISCLAIMER)
    run.italic = True
    run.font.size = Pt(8)
    doc.sections[0].footer.paragraphs[0].text = config.COMPANY_FOOTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
def build_pdf(title, body, terms, include_logo=True) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    text_all = title + body + "".join(k + v for k, v in terms)
    font, bold = "Helvetica", "Helvetica-Bold"
    if custom_font_available():
        pdfmetrics.registerFont(TTFont("LE-Font", str(config.FONT_PATH)))
        font = bold = "LE-Font"
        if Path(config.FONT_BOLD_PATH).exists():
            pdfmetrics.registerFont(TTFont("LE-Bold", str(config.FONT_BOLD_PATH)))
            bold = "LE-Bold"
    elif _has_non_latin(text_all):
        raise ValueError(
            "This text needs a Unicode font for PDF. Put a .ttf file at fonts/pdf_font.ttf "
            "(for example Noto Sans Tamil / Devanagari), or export DOCX/TXT instead."
        )

    s_title = ParagraphStyle("t", fontName=bold, fontSize=18, leading=22, alignment=TA_CENTER, spaceAfter=10)
    s_h2 = ParagraphStyle("h2", fontName=bold, fontSize=12, leading=15, spaceBefore=10, spaceAfter=4)
    s_body = ParagraphStyle("b", fontName=font, fontSize=10.5, leading=14.5, spaceAfter=5)
    s_bul = ParagraphStyle("bl", parent=s_body, leftIndent=14, bulletIndent=4)
    s_small = ParagraphStyle("s", fontName=font, fontSize=8, leading=10, textColor=colors.grey)

    story = []
    if include_logo and Path(config.LOGO_PATH).exists():
        story.append(Image(str(config.LOGO_PATH), width=45 * mm, height=15 * mm, kind="proportional", hAlign="LEFT"))
        story.append(Spacer(1, 6))
    story.append(Paragraph(xml_escape(title), s_title))

    if terms:
        story.append(Paragraph("Key Terms", s_h2))
        rows = [[Paragraph(xml_escape(k), ParagraphStyle("k", fontName=bold, fontSize=9.5, leading=12)),
                 Paragraph(xml_escape(v), ParagraphStyle("v", fontName=font, fontSize=9.5, leading=12))]
                for k, v in terms]
        t = Table(rows, colWidths=[55 * mm, 115 * mm])
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF1F6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story += [t, Spacer(1, 8)]

    for kind, text in parse_blocks(body):
        if kind == "h1":
            continue
        if kind == "h2":
            story.append(Paragraph(xml_escape(text), s_h2))
        elif kind == "bullet":
            story.append(Paragraph(xml_escape(text), s_bul, bulletText="•"))
        else:
            story.append(Paragraph(xml_escape(text), s_body))
    story += [Spacer(1, 10), Paragraph(xml_escape(config.DISCLAIMER), s_small)]

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(20 * mm, 10 * mm, config.COMPANY_FOOTER)
        canvas.drawRightString(190 * mm, 10 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                      topMargin=18 * mm, bottomMargin=20 * mm, title=title).build(
        story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
